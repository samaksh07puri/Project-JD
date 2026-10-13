"""Render a structured resume to PDF (ReportLab) and DOCX (python-docx)."""
from __future__ import annotations

import io
import re
from html import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from ..schemas import ExportRequest

ACCENT = HexColor("#1d4ed8")


def _safe(text: str) -> str:
    return escape(text or "").strip()


def _paragraphs(text: str) -> list[str]:
    """Split a plain-text document into trimmed, non-empty paragraphs."""
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        return []
    parts = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    if len(parts) == 1:
        parts = [part.strip() for part in normalized.split("\n") if part.strip()]
    return parts


def render_pdf(resume: ExportRequest) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=f"{resume.contact.name or 'Resume'} - Tailored Resume",
    )

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name", parent=styles["Title"], fontSize=22, textColor=HexColor("#0f172a"),
        spaceAfter=2, alignment=TA_CENTER,
    )
    headline_style = ParagraphStyle(
        "Headline", parent=styles["Normal"], fontSize=11.5, textColor=ACCENT,
        spaceAfter=2, alignment=TA_CENTER,
    )
    contact_style = ParagraphStyle(
        "Contact", parent=styles["Normal"], fontSize=9, textColor=HexColor("#475569"),
        spaceAfter=10, alignment=TA_CENTER,
    )
    section_style = ParagraphStyle(
        "Section", parent=styles["Heading2"], fontSize=12, textColor=ACCENT,
        spaceBefore=12, spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontSize=10, leading=14, textColor=HexColor("#111827"),
    )
    sub_style = ParagraphStyle(
        "Sub", parent=styles["Normal"], fontSize=10.5, leading=14,
        textColor=HexColor("#0f172a"),
    )
    muted_style = ParagraphStyle(
        "Muted", parent=styles["Normal"], fontSize=9, textColor=HexColor("#64748b"),
    )

    story: list = []

    if resume.contact.name:
        story.append(Paragraph(_safe(resume.contact.name), name_style))
    if resume.contact.headline:
        story.append(Paragraph(_safe(resume.contact.headline), headline_style))
    contact_bits = [
        resume.contact.email,
        resume.contact.phone,
        resume.contact.location,
        *resume.contact.links,
    ]
    contact_line = "  -  ".join(bit for bit in contact_bits if bit)
    if contact_line:
        story.append(Paragraph(_safe(contact_line), contact_style))

    def add_section(title: str) -> None:
        story.append(Paragraph(_safe(title.upper()), section_style))

    if resume.summary:
        add_section("Professional Summary")
        story.append(Paragraph(_safe(resume.summary), body_style))

    if resume.skills:
        add_section("Skills")
        story.append(Paragraph(_safe(" - ".join(resume.skills)), body_style))

    if resume.experience:
        add_section("Experience")
        for job in resume.experience:
            header = " | ".join(part for part in [job.title, job.company] if part)
            story.append(Paragraph(f"<b>{_safe(header)}</b>", sub_style))
            meta = " | ".join(part for part in [job.location, job.dates] if part)
            if meta:
                story.append(Paragraph(_safe(meta), muted_style))
            if job.bullets:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(_safe(b), body_style), leftIndent=14) for b in job.bullets],
                        bulletType="bullet", start="-", leftIndent=12, bulletFontSize=7,
                    )
                )
            story.append(Spacer(1, 6))

    if resume.projects:
        add_section("Projects")
        for project in resume.projects:
            story.append(Paragraph(f"<b>{_safe(project.name)}</b>", sub_style))
            if project.description:
                story.append(Paragraph(_safe(project.description), body_style))
            if project.bullets:
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(_safe(b), body_style), leftIndent=14) for b in project.bullets],
                        bulletType="bullet", start="-", leftIndent=12, bulletFontSize=7,
                    )
                )
            story.append(Spacer(1, 6))

    if resume.education:
        add_section("Education")
        for edu in resume.education:
            line = " | ".join(part for part in [edu.degree, edu.institution] if part)
            story.append(Paragraph(f"<b>{_safe(line)}</b>", sub_style))
            meta = " | ".join(part for part in [edu.location, edu.dates] if part)
            if meta:
                story.append(Paragraph(_safe(meta), muted_style))
            if edu.details:
                story.append(Paragraph(_safe(edu.details), body_style))
            story.append(Spacer(1, 4))

    if resume.certifications:
        add_section("Certifications")
        for cert in resume.certifications:
            story.append(Paragraph(f"- {_safe(cert)}", body_style))

    doc.build(story)
    return buffer.getvalue()


def render_docx(resume: ExportRequest) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def heading(text: str) -> None:
        p = document.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)
        p.space_before = Pt(10)
        p.space_after = Pt(2)

    if resume.contact.name:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(resume.contact.name)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    if resume.contact.headline:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(resume.contact.headline)
        run.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

    contact_bits = [
        resume.contact.email, resume.contact.phone, resume.contact.location,
        *resume.contact.links,
    ]
    contact_line = "  -  ".join(bit for bit in contact_bits if bit)
    if contact_line:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(contact_line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    if resume.summary:
        heading("Professional Summary")
        document.add_paragraph(resume.summary)

    if resume.skills:
        heading("Skills")
        document.add_paragraph(" - ".join(resume.skills))

    if resume.experience:
        heading("Experience")
        for job in resume.experience:
            header = " | ".join(part for part in [job.title, job.company] if part)
            p = document.add_paragraph()
            run = p.add_run(header)
            run.bold = True
            meta = " | ".join(part for part in [job.location, job.dates] if part)
            if meta:
                p = document.add_paragraph()
                run = p.add_run(meta)
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            for bullet in job.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    if resume.projects:
        heading("Projects")
        for project in resume.projects:
            p = document.add_paragraph()
            run = p.add_run(project.name)
            run.bold = True
            if project.description:
                document.add_paragraph(project.description)
            for bullet in project.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    if resume.education:
        heading("Education")
        for edu in resume.education:
            line = " | ".join(part for part in [edu.degree, edu.institution] if part)
            p = document.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            meta = " | ".join(part for part in [edu.location, edu.dates] if part)
            if meta:
                document.add_paragraph(meta)
            if edu.details:
                document.add_paragraph(edu.details)

    if resume.certifications:
        heading("Certifications")
        for cert in resume.certifications:
            document.add_paragraph(cert, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_cover_letter_pdf(text: str, name: str = "") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=0.9 * inch,
        rightMargin=0.9 * inch,
        topMargin=0.9 * inch,
        bottomMargin=0.9 * inch,
        title="Cover Letter",
    )

    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "CLName", parent=styles["Title"], fontSize=18, textColor=HexColor("#0f172a"),
        alignment=0, spaceAfter=14,
    )
    body_style = ParagraphStyle(
        "CLBody", parent=styles["Normal"], fontSize=11, leading=16,
        textColor=HexColor("#111827"), spaceAfter=10,
    )

    story: list = []
    if name:
        story.append(Paragraph(_safe(name), name_style))
    for paragraph in _paragraphs(text):
        story.append(Paragraph(_safe(paragraph), body_style))

    doc.build(story)
    return buffer.getvalue()


def render_cover_letter_docx(text: str, name: str = "") -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    if name:
        p = document.add_paragraph()
        run = p.add_run(name)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for paragraph in _paragraphs(text):
        document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
