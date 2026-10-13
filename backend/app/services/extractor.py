"""Extract plain text from uploaded PDF and DOCX resumes."""
from __future__ import annotations

import io
import re

import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = (".pdf", ".docx")

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


class ExtractionError(Exception):
    """Raised when a resume cannot be parsed into text."""


def _clean(text: str) -> str:
    """Normalise whitespace while keeping paragraph/line structure."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(data: bytes) -> str:
    pages: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(page_text)
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    """Return cleaned text extracted from ``data`` based on the file extension.

    Raises:
        ExtractionError: for unsupported file types or unparsable documents.
    """
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            text = _extract_pdf(data)
        elif name.endswith(".docx"):
            text = _extract_docx(data)
        else:
            raise ExtractionError(
                "Unsupported file type. Please upload a PDF or DOCX file."
            )
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - convert any parser error to a clean message
        raise ExtractionError(
            f"Could not read the resume file. It may be corrupted or password protected. ({exc})"
        ) from exc

    cleaned = _clean(text)
    if len(cleaned) < 50:
        raise ExtractionError(
            "We could not extract readable text from this file. "
            "If it is a scanned image PDF, please upload a text-based PDF or DOCX."
        )
    return cleaned
